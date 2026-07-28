"""Regenera KOL_Tracker.xlsx (informe) a partir de los CSVs de data/.

Los CSVs son la fuente de verdad; este Excel es un informe generado.
Las hojas KOLs/Posts/Pagos llevan los datos como valores y las columnas
de resumen como formulas, para que el fichero siga recalculando bien
si se abre en Excel o se importa a Google Sheets.
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
DATA_FONT = Font(name=ARIAL, size=10, color="0000FF")  # azul: datos
F_FONT = Font(name=ARIAL, size=10)                     # negro: formulas
TITLE_FONT = Font(name=ARIAL, bold=True, size=14, color="1F3864")
SUB_FONT = Font(name=ARIAL, bold=True, size=11, color="1F3864")
NOTE_FONT = Font(name=ARIAL, size=9, italic=True, color="808080")
THIN = Border(*[Side(style="thin", color="B0B0B0")] * 4)

MONEY = '$#,##0.00;($#,##0.00);"-"'
MONEY0 = '$#,##0;($#,##0);"-"'
NUM = '#,##0;(#,##0);"-"'


def _headers(ws, cols, widths):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def _num(v):
    try:
        return float(str(v).replace(",", "."))
    except (TypeError, ValueError):
        return None


def build(kols, posts, pagos, out_path):
    maxrow = max(300, len(kols) + 50, len(posts) + 50, len(pagos) + 50)
    wb = Workbook()

    # ---- Guia ----
    g = wb.active
    g.title = "Guía"
    g.column_dimensions["A"].width = 3
    g.column_dimensions["B"].width = 110
    lines = [
        (TITLE_FONT, "KOL Tracker — GenLayer Ambassadors & Rally"),
        (None, ""),
        (SUB_FONT, "Este Excel es un informe generado automáticamente"),
        (None, "Los datos viven en kol-tracker/data/*.csv y el script de automatización regenera este fichero."),
        (None, "No edites este Excel: los cambios se perderán en la siguiente ejecución. Edita los CSVs."),
        (None, ""),
        (SUB_FONT, "Qué hace la automatización"),
        (None, "• Vigila el timeline de cada KOL activo y añade los posts que mencionan la campaña."),
        (None, "• Refresca views/likes/RTs de los posts recientes."),
        (None, "• Genera cada mes las filas de pago (fijos mensuales y devengo por post)."),
        (None, "• Tú solo cambias el estado de un pago de 'Pendiente' a 'Pagado' en data/pagos.csv."),
        (None, ""),
        (SUB_FONT, "Código de colores"),
        (None, "• Texto azul = datos. • Texto negro = fórmulas (Dashboard y columnas de resumen)."),
    ]
    for r, (font, txt) in enumerate(lines, start=1):
        g.cell(row=r, column=2, value=txt).font = font or Font(name=ARIAL, size=10)

    # ---- KOLs ----
    k = wb.create_sheet("KOLs")
    _headers(k, ["Nombre", "Handle X", "Link X", "Programa", "Tipo de acuerdo", "Tarifa ($)",
                 "Fecha inicio", "Estado", "Wallet / método de pago", "Notas",
                 "Nº posts", "Views totales", "Total pagado ($)", "Pendiente ($)", "CPM ($/1k views)"],
              [20, 16, 30, 20, 16, 11, 12, 10, 24, 26, 9, 12, 13, 12, 13])
    for i, row in enumerate(kols, start=2):
        vals = [row.get(f, "") for f in ("nombre", "handle", "link", "programa", "tipo_acuerdo")]
        vals += [_num(row.get("tarifa"))] + [row.get(f, "") for f in ("fecha_inicio", "estado", "pago_metodo", "notas")]
        for j, v in enumerate(vals, start=1):
            c = k.cell(row=i, column=j, value=v)
            c.font = DATA_FONT
            c.border = THIN
        k.cell(row=i, column=6).number_format = MONEY0
    for i in range(2, maxrow + 1):
        a = f"A{i}"
        formulas = [
            f'=IF({a}="","",COUNTIF(Posts!$B$2:$B${maxrow},{a}))',
            f'=IF({a}="","",SUMIF(Posts!$B$2:$B${maxrow},{a},Posts!$F$2:$F${maxrow}))',
            f'=IF({a}="","",SUMIFS(Pagos!$D$2:$D${maxrow},Pagos!$B$2:$B${maxrow},{a},Pagos!$F$2:$F${maxrow},"Pagado"))',
            f'=IF({a}="","",SUMIFS(Pagos!$D$2:$D${maxrow},Pagos!$B$2:$B${maxrow},{a},Pagos!$F$2:$F${maxrow},"Pendiente"))',
            f'=IF({a}="","",IF(L{i}=0,"-",M{i}/L{i}*1000))',
        ]
        for j, (f, fmt) in enumerate(zip(formulas, [NUM, NUM, MONEY0, MONEY0, MONEY]), start=11):
            c = k.cell(row=i, column=j, value=f)
            c.font = F_FONT
            c.number_format = fmt
            c.border = THIN

    # ---- Posts ----
    p = wb.create_sheet("Posts")
    _headers(p, ["Fecha", "KOL", "Programa (auto)", "Link al post", "Tipo", "Views", "Likes",
                 "RTs", "Comentarios", "Coste atribuido ($)", "Notas"],
             [12, 20, 20, 42, 10, 11, 9, 9, 12, 15, 26])
    for i, row in enumerate(posts, start=2):
        vals = [row.get("fecha", ""), row.get("kol", ""), None, row.get("link", ""), row.get("tipo", ""),
                _num(row.get("views")), _num(row.get("likes")), _num(row.get("rts")),
                _num(row.get("replies")), _num(row.get("coste")), row.get("notas", "")]
        for j, v in enumerate(vals, start=1):
            if j == 3:
                continue
            c = p.cell(row=i, column=j, value=v)
            c.font = DATA_FONT
            c.border = THIN
        for j in (6, 7, 8, 9):
            p.cell(row=i, column=j).number_format = NUM
        p.cell(row=i, column=10).number_format = MONEY0
    for i in range(2, maxrow + 1):
        c = p.cell(row=i, column=3, value=(
            f'=IF(B{i}="","",IFERROR(INDEX(KOLs!$D$2:$D${maxrow},'
            f'MATCH(B{i},KOLs!$A$2:$A${maxrow},0)),"¿KOL no dado de alta?"))'))
        c.font = F_FONT
        c.border = THIN

    # ---- Pagos ----
    pg = wb.create_sheet("Pagos")
    _headers(pg, ["Fecha", "KOL", "Programa (auto)", "Importe ($)", "Concepto", "Estado",
                  "Método / tx", "Notas"],
             [12, 20, 20, 12, 30, 12, 26, 26])
    for i, row in enumerate(pagos, start=2):
        vals = [row.get("fecha", ""), row.get("kol", ""), None, _num(row.get("importe")),
                row.get("concepto", ""), row.get("estado", ""), row.get("metodo", ""), row.get("notas", "")]
        for j, v in enumerate(vals, start=1):
            if j == 3:
                continue
            c = pg.cell(row=i, column=j, value=v)
            c.font = DATA_FONT
            c.border = THIN
        pg.cell(row=i, column=4).number_format = MONEY0
    for i in range(2, maxrow + 1):
        c = pg.cell(row=i, column=3, value=(
            f'=IF(B{i}="","",IFERROR(INDEX(KOLs!$D$2:$D${maxrow},'
            f'MATCH(B{i},KOLs!$A$2:$A${maxrow},0)),"¿KOL no dado de alta?"))'))
        c.font = F_FONT
        c.border = THIN

    # ---- Dashboard ----
    d = wb.create_sheet("Dashboard", 1)
    for col, w in zip("ABCDEFGH", [30, 14, 14, 16, 16, 14, 15, 15]):
        d.column_dimensions[col].width = w
    d["A1"] = "KOL Tracker — Dashboard"
    d["A1"].font = TITLE_FONT
    d["A2"] = "Generado automáticamente a partir de kol-tracker/data/*.csv."
    d["A2"].font = NOTE_FONT
    d["A4"] = "Totales"
    d["A4"].font = SUB_FONT
    tot = [
        ("KOLs activos", f'=COUNTIF(KOLs!$H$2:$H${maxrow},"Activo")', NUM),
        ("Posts publicados", f'=COUNTA(Posts!$A$2:$A${maxrow})', NUM),
        ("Views totales", f'=SUM(Posts!$F$2:$F${maxrow})', NUM),
        ("Total pagado ($)", f'=SUMIF(Pagos!$F$2:$F${maxrow},"Pagado",Pagos!$D$2:$D${maxrow})', MONEY0),
        ("Total pendiente ($)", f'=SUMIF(Pagos!$F$2:$F${maxrow},"Pendiente",Pagos!$D$2:$D${maxrow})', MONEY0),
        ("Coste por post ($)", '=IF(B6=0,"-",B8/B6)', MONEY),
        ("CPM ($/1k views)", '=IF(B7=0,"-",B8/B7*1000)', MONEY),
    ]
    for i, (label, f, fmt) in enumerate(tot, start=5):
        d.cell(row=i, column=1, value=label).font = F_FONT
        c = d.cell(row=i, column=2, value=f)
        c.font = Font(name=ARIAL, size=10, bold=True)
        c.number_format = fmt
        c.border = THIN
    d["A14"] = "Por programa"
    d["A14"].font = SUB_FONT
    for j, h in enumerate(["Programa", "KOLs activos", "Posts", "Views", "Pagado ($)",
                           "Pendiente ($)", "Coste/post ($)", "CPM ($/1k)"], start=1):
        c = d.cell(row=15, column=j, value=h)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = THIN
    for i, prog in enumerate(["GenLayer Ambassador", "Rally"], start=16):
        d.cell(row=i, column=1, value=prog).font = F_FONT
        d.cell(row=i, column=1).border = THIN
        formulas = [
            f'=COUNTIFS(KOLs!$D$2:$D${maxrow},$A{i},KOLs!$H$2:$H${maxrow},"Activo")',
            f'=COUNTIF(Posts!$C$2:$C${maxrow},$A{i})',
            f'=SUMIF(Posts!$C$2:$C${maxrow},$A{i},Posts!$F$2:$F${maxrow})',
            f'=SUMIFS(Pagos!$D$2:$D${maxrow},Pagos!$C$2:$C${maxrow},$A{i},Pagos!$F$2:$F${maxrow},"Pagado")',
            f'=SUMIFS(Pagos!$D$2:$D${maxrow},Pagos!$C$2:$C${maxrow},$A{i},Pagos!$F$2:$F${maxrow},"Pendiente")',
            f'=IF(C{i}=0,"-",E{i}/C{i})',
            f'=IF(D{i}=0,"-",E{i}/D{i}*1000)',
        ]
        for j, (f, fmt) in enumerate(zip(formulas, [NUM, NUM, NUM, MONEY0, MONEY0, MONEY, MONEY]), start=2):
            c = d.cell(row=i, column=j, value=f)
            c.font = F_FONT
            c.number_format = fmt
            c.border = THIN
    d["A20"] = "El detalle por KOL (posts, pagado, pendiente, CPM) está en la hoja 'KOLs'."
    d["A20"].font = NOTE_FONT

    wb.save(out_path)
