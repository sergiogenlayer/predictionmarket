#!/usr/bin/env python3
"""Permissionless Campaigns activation calendar — announcements & activations only,
X-first (no Discord/Telegram), sales-to-projects focus, Mailing for key articles/calls."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/tmp/claude-0/-home-user-predictionmarket/cadac25f-7355-5a2e-945b-6583f4b657ab/scratchpad/permissionless_campaigns_calendar_v2.xlsx"

FONT = "Arial"
thin = Side(style="thin", color="B7B7B7")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_white_bold = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=12, bold=True, color="FFFFFF")

fill_week = PatternFill("solid", fgColor="134F5C")
fill_days = PatternFill("solid", fgColor="D9D9D9")
fill_cat = PatternFill("solid", fgColor="EFEFEF")
fill_action = PatternFill("solid", fgColor="FCE5CD")   # action release (orange)
fill_video = PatternFill("solid", fgColor="CFE2F3")    # video pieces (blue)
fill_annc = PatternFill("solid", fgColor="D9EAD3")     # X announcements (green)
fill_article = PatternFill("solid", fgColor="FFF2CC")  # articles (yellow)
fill_sales = PatternFill("solid", fgColor="EAD1DC")    # sales-to-projects pieces (pink)
fill_title = PatternFill("solid", fgColor="0B3040")

wrap = Alignment(wrap_text=True, vertical="top", horizontal="left")
center = Alignment(wrap_text=True, vertical="center", horizontal="center")

wb = Workbook()
ws = wb.active
ws.title = "Permissionless Campaigns"

ws.column_dimensions["A"].width = 22
for c in range(2, 9):
    ws.column_dimensions[get_column_letter(c)].width = 27

ws.merge_cells("A1:H1")
ws["A1"] = "PERMISSIONLESS CAMPAIGNS — ACTIVATION CALENDAR · Aug 3 – Sep 6, 2026 · X-first, sell to projects"
ws["A1"].font = f_title
ws["A1"].fill = fill_title
ws["A1"].alignment = center
ws.row_dimensions[1].height = 28

days = ["Category", "Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
for i, d in enumerate(days, start=1):
    c = ws.cell(row=2, column=i, value=d)
    c.font = f_bold
    c.fill = fill_days
    c.alignment = center
    c.border = border
ws.freeze_panes = "B3"

CHANNELS = ["Action Release", "Video", "X — Announcement", "X — Sell to Projects",
            "Article", "X Space / Live", "Creators & Communities", "BD / Projects",
            "PR", "Mailing", "KOL"]

FILLS = {"Action Release": fill_action, "Video": fill_video,
         "X — Announcement": fill_annc, "Article": fill_article,
         "X — Sell to Projects": fill_sales}

weeks = [
    (
        "Week of Aug 03 – Aug 09, 2026 (NARRATIVE + TEASER)",
        {
            "Video": ["", "", "", "", "Demo teaser — 15s cutdown of the Demo Video", "", ""],
            "X — Announcement": ["", "", "", "", "Something permissionless is coming (TEASER POST)", "", "Countdown — launch date reveal (POST)"],
            "Article": ["", "The Protocol Narrative — campaigns without gatekeepers (ARTICLE)", "", "", "", "", ""],
            "Creators & Communities": ["", "", "Send assets/copy kit to communities & creators (private)", "Brief 3-4 creators for testimonial videos (record W2-W3)", "", "", ""],
            "BD / Projects": ["", "", "Build target list of projects launching in Sept-Oct (pipeline)", "Close flagship community for co-created Network Campaign (launches W2)", "", "", ""],
        },
    ),
    (
        "Week of Aug 10 – Aug 16, 2026 (LAUNCH)",
        {
            "Action Release": ["PERMISSIONLESS CAMPAIGNS LIVE", "", "", "", "", "", ""],
            "Video": ["DEMO VIDEO — hero piece: create & launch a campaign end-to-end", "", "", "", "", "", ""],
            "X — Announcement": ["Launch thread + Demo Video (pinned)", "", "", "", "", "", ""],
            "X — Sell to Projects": [
                "", "",
                "Verified Campaigns — what 1,500$ gets your project: protocol visibility, Rally marketing pack, elite creators & ambassadors (THREAD + ONE-PAGER)",
                "Network Campaigns — permissionless, rewards in USDC | RLPs (THREAD + Verified vs Network VISUAL)",
                "CALL FOR PROJECTS — launching soon? Run your launch as a Network Campaign (POST + form)",
                "", "",
            ],
            "Article": ["", "Permissionless Campaigns: how it works — pay in RLPs | USDC (ARTICLE)", "", "", "", "", ""],
            "Creators & Communities": ["Coordinated push — communities & creators RT/QT the launch simultaneously", "", "", "", "Flagship co-created Network Campaign goes live (partner community)", "", ""],
            "PR": ["", "", "", "", "", "", ""],
            "Mailing": ["Launch newsletter", "", "", "Verified one-pager to project pipeline", "", "", ""],
            "KOL": ["Amplification", "", "", "", "", "", ""],
        },
    ),
    (
        "Week of Aug 17 – Aug 23, 2026 (SELL TO PROJECTS + CREATORS)",
        {
            "Video": ["", "", "Creator video #1 — 'I launched a campaign, here's my experience'", "", "Demo cutdown #2 — project POV: 'your launch could look like this'", "", ""],
            "X — Announcement": ["First 10 Network Campaigns — spotlight (THREAD)", "", "", "", "", "", ""],
            "X — Sell to Projects": ["", "Why your project should create a Network Campaign — 5 use cases (THREAD, links article)", "", "", "", "", ""],
            "Article": ["", "Why create a Network Campaign — 5 use cases (ARTICLE, sales-ready)", "", "Viraliza tu Launch — playbook for projects (ARTICLE)", "", "", ""],
            "X Space / Live": ["", "", "Live demo for projects — launch a campaign in 10 minutes (X SPACE/VIDEO)", "", "", "X Space with communities & creators — first campaigns debrief", ""],
            "Creators & Communities": ["", "", "Ambassadors QT the demo with their own take", "", "", "", ""],
            "BD / Projects": ["1:1 outreach to projects with Verified one-pager (all week)", "", "", "", "", "", ""],
            "Mailing": ["", "", "", "'Why Network Campaigns' article to pipeline + book-your-campaign call", "", "", ""],
            "KOL": ["", "", "Amplification", "", "", "", ""],
        },
    ),
    (
        "Week of Aug 24 – Aug 30, 2026 (SOCIAL PROOF + PR)",
        {
            "Video": ["", "", "Creator video #2 — community lead POV", "", "", "", ""],
            "X — Announcement": ["Campaign of the Week — spotlight (POST, recurring)", "", "", "", "Best campaigns so far — curated recap (POST)", "", ""],
            "X — Sell to Projects": ["", "", "", "Case study THREAD: first Verified Campaign — results, ROI, what they unlocked", "", "", ""],
            "Article": ["", "", "", "Case study: first Verified Campaign (ARTICLE)", "", "", ""],
            "BD / Projects": ["Verified pipeline follow-ups (use case study as proof)", "", "", "", "", "", ""],
            "PR": ["", "Press release with first numbers (campaigns created, rewards paid)", "", "", "", "", ""],
            "Mailing": ["", "", "", "Case study to pipeline — 'this could be your launch'", "", "", ""],
            "KOL": ["", "Amplification of PR", "", "", "", "", ""],
        },
    ),
    (
        "Week of Aug 31 – Sep 06, 2026 (MOMENTUM — light week)",
        {
            "Video": ["", "", "Creator video #3 / month recap reel (cutdowns compilation)", "", "", "", ""],
            "X — Announcement": ["Numbers post — month in review (show numbers) + Campaign of the Week", "", "", "", "", "", "What's next — roadmap tease (POST)"],
            "X — Sell to Projects": ["", "", "", "", "Open call: September cohort for Network & Verified Campaigns (POST)", "", ""],
            "Mailing": ["", "", "", "", "Monthly recap + September cohort open call", "", ""],
        },
    ),
]

row = 3
for banner, data in weeks:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    b = ws.cell(row=row, column=1, value=banner)
    b.font = f_white_bold
    b.fill = fill_week
    b.alignment = center
    for c in range(1, 9):
        ws.cell(row=row, column=c).border = border
    ws.row_dimensions[row].height = 22
    row += 1

    for ch in CHANNELS:
        cat = ws.cell(row=row, column=1, value=ch)
        cat.font = f_bold
        cat.fill = fill_cat
        cat.alignment = wrap
        cat.border = border
        vals = data.get(ch, [""] * 7)
        has_content = any(vals)
        for i, v in enumerate(vals):
            cell = ws.cell(row=row, column=2 + i, value=v)
            cell.font = f_base
            cell.alignment = wrap
            cell.border = border
            if v and ch in FILLS:
                cell.fill = FILLS[ch]
                if ch == "Action Release":
                    cell.font = f_bold
        ws.row_dimensions[row].height = 50 if has_content else 16
        row += 1
    row += 1

notes = [
    ("How to read this calendar", ""),
    ("Logic of the month", "W1 narrative + teaser -> W2 launch (Demo Video + Verified/Network threads + call for projects) -> W3 sell to projects + creators (live demo + Space ~1 week after launch) -> W4 social proof + PR (~2 weeks after launch, with real numbers) -> W5 momentum, light."),
    ("X-first", "Everything runs on X (posts, threads, Spaces, live demo). No Discord/Telegram. 'X — Sell to Projects' (pink) is content aimed directly at projects: Verified/Network pitches, call for projects, case study ROI."),
    ("Mailing", "Only for key articles and important calls to the project pipeline: launch, Verified one-pager, 'Why Network Campaigns', case study, monthly recap + cohort call."),
    ("Reusable assets", "Demo Video (+ cutdowns), Verified one-pager, Verified vs Network visual, use-cases & playbook articles, creator videos and the case study all double as sales material for pitching projects."),
    ("Verified Campaigns", "Verification filter + 1,500$. They get: protocol visibility/positioning, marketing service pack from Rally socials (articles, videos), and unlock of community & elite creators features (ambassadors + campaign management in Rally)."),
    ("Network Campaigns", "Permissionless — anyone can create a campaign. Rewards only in USDC or RLPs."),
    ("Creator videos", "3 short testimonial videos (briefed W1, recorded W2-W3, published W3-W5): creator experience, community lead POV, month recap."),
    ("", ""),
    ("Internal Campaigns doc", "https://docs.google.com/document/d/13DhJE-HBozbZijic54TgX2L6m8FPw_oHp-_uZ1aOzuQ/edit?usp=sharing"),
    ("Rally Tracker", "https://rally-tracker-ten.vercel.app/"),
]
for label, text in notes:
    a = ws.cell(row=row, column=1, value=label)
    a.font = f_bold
    a.alignment = wrap
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
    b = ws.cell(row=row, column=2, value=text)
    b.font = f_base
    b.alignment = wrap
    if text:
        ws.row_dimensions[row].height = 30
    row += 1

wb.save(OUT)
print("saved", OUT)
